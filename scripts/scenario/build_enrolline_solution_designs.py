#!/usr/bin/env python3
"""Build the two supplied Enrolline Solution Designs (.docx) — in-world documents.

Enrolline is the CL3 *assessment* system: YAT's student records and enrolment management system,
sitting on a single-AZ cloud baseline that CL3 improves. Two documents come out of this one
generator, because they are two states of the same design and must not drift apart:

  BASELINE  — "Enrolline Cloud Architecture — Baseline Design", the as-built single-AZ state.
              The current-state record AT1 analyses. Visible from s1-cl3-at1.
  IMPROVED  — "Enrolline Cloud Architecture — Approved Improvement Design", the agreed "to be"
              design handed to the team as the AT2 build input. Visible from s1-cl3-at2.

Parallel-but-different against the Ledgerline (CL1 practice) baseline: 10.30.0.0/16 (not 10.20),
an S3 document store that is a working part of the system (Ledgerline has none), an intake/census
load profile rather than a month-end one, and 30-year student-record retention rather than 7-year
financial retention. In-world artefacts — no UoC tags.

Usage:  python scripts/scenario/build_enrolline_solution_designs.py [--only baseline|improved]
Output: ../diploma-cloud-cyber-website-s1/public/documents/YAT-Enrolline-{Baseline,Improved}-Solution-Design.docx
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # content-repo scripts/ (brand + registry)  # noqa: E402
sys.path.insert(0, str(next(d / "scripts" for d in Path(__file__).resolve().parents if (d / "scripts" / "helpers" / "__init__.py").exists())))  # umbrella scripts/ (engine)  # noqa: E402
from helpers.docx_body_text import add_body_paragraph, add_bullet_list  # noqa: E402
from helpers.docx_tables import add_data_table  # noqa: E402
from helpers.docx_styling import add_field, paragraph_bottom_rule, set_cell_borders, shade_cell  # noqa: E402
from brand import ADDRESS, CREAM, GREY, TEAL, TERRACOTTA  # noqa: E402
from helpers.scenario_document import build_header_footer, configure_styles, wordmark  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.shared import Pt, Cm, RGBColor  # noqa: E402

OUT_DIR = (Path(__file__).resolve().parents[3] / "diploma-cloud-cyber-website-s1"
           / "public" / "documents")

# Published diagrams: this topology is also served as an intranet page, so the .drawio/.png live at
# the top level and one asset serves both surfaces. The .png is rendered by the draw-diagram skill
# from scripts/scenario/diagrams/network-enrolline-baseline-singleaz.json — no hand export.
DIAGRAM_DIR = (Path(__file__).resolve().parents[3] / "diploma-cloud-cyber-website-s1"
               / "public" / "diagrams")


def diagram_figure(doc, caption, image_name, width_cm=16.0):
    """Place a generated diagram, captioned.

    The generator places the picture, so it survives a rebuild. Fails loudly if the image is
    absent rather than leaving a silent hole in the document.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    image = DIAGRAM_DIR / image_name
    if not image.exists():
        raise FileNotFoundError(
            f"Diagram not found: {image}\n"
            f"Render it from its spec: .claude/skills/draw-diagram/.venv/bin/python "
            f".claude/skills/draw-diagram/draw_diagram.py --spec "
            f"scripts/scenario/diagrams/network-enrolline-baseline-singleaz.json "
            f"--out <…>.drawio --png <…>.png")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image), width=Cm(width_cm))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.italic = True; cr.font.size = Pt(9); cr.font.color.rgb = RGBColor.from_string(GREY)
    return p


def na(doc, reason):
    p = doc.add_paragraph()
    r = p.add_run(f"Not applicable — {reason}")
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(TERRACOTTA)
    p.paragraph_format.space_after = Pt(6)
    return p


def lab_note(doc, text):
    """Where the design and the build environment differ, say so in one plain sentence."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    lead = p.add_run("Building this in a lab ")
    lead.bold = True; lead.font.size = Pt(9.5)
    lead.font.color.rgb = RGBColor.from_string(TERRACOTTA)
    r = p.add_run(text)
    r.italic = True; r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(TERRACOTTA)
    return p


def page_setup(doc):
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.6); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)
    build_header_footer(sec)


def cover(doc, subtitle, rows):
    wordmark(doc.add_paragraph())
    ar = doc.add_paragraph().add_run(ADDRESS)
    ar.font.size = Pt(9); ar.font.color.rgb = RGBColor.from_string(GREY)
    paragraph_bottom_rule(doc.add_paragraph(), TEAL, sz=12)
    for _ in range(3):
        doc.add_paragraph()
    doc.add_paragraph(style="Title").add_run("Solution Design")
    sub = doc.add_paragraph().add_run(subtitle)
    sub.font.size = Pt(15); sub.bold = True; sub.font.color.rgb = RGBColor.from_string(TERRACOTTA)
    doc.add_paragraph()
    ct = doc.add_table(rows=0, cols=2)
    for k, v in rows:
        cells = ct.add_row().cells
        set_cell_borders(cells[0]); set_cell_borders(cells[1]); shade_cell(cells[0], CREAM)
        kr = cells[0].paragraphs[0].add_run(k); kr.bold = True; kr.font.size = Pt(10)
        cells[1].paragraphs[0].add_run(v).font.size = Pt(10)
        cells[0].width = Cm(4.5); cells[1].width = Cm(12.0)


def contents(doc):
    doc.add_section(WD_SECTION.NEW_PAGE); build_header_footer(doc.sections[-1])
    doc.add_paragraph("Contents", style="Heading 1")
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
              placeholder="Right-click and choose “Update Field” to build the table of contents.")
    doc.add_section(WD_SECTION.NEW_PAGE); build_header_footer(doc.sections[-1])


def document_control(doc, h1, version, approved):
    h1("Document control")
    add_data_table(doc, ["Field", "Value"], [
        ["Document owner", "Pat Lin, MTS Senior Consultant"],
        ["Business owner", "Jess Tran, YAT Registrar"],
        ["Version", version],
        ["Approved by", approved],
        ["Review cycle", "On material change to the architecture or its operating profile"],
        ["Classification", "Internal — YAT ICT, Student Administration, Compliance, and MTS personnel on signed MSA"],
    ], widths=[5.0, 11.5])


# --------------------------------------------------------------------------------------
# BASELINE — the as-built single-AZ state
# --------------------------------------------------------------------------------------

def build_baseline(path):
    doc = Document()
    configure_styles(doc)
    page_setup(doc)

    cover(doc, "YAT Enrolline Cloud Architecture — Baseline Design", [
        ("Engagement", "YAT Enrolline Cloud Migration — Foundation Build"),
        ("Document type", "Technical design (Solution Design)"),
        ("Version", "v1.0 — Approved for implementation"),
        ("Authored by", "MTS Senior Architecture Team, in consultation with YAT ICT and the YAT Registrar"),
        ("Approved by", "Pat Lin (MTS Senior Consultant) · Sam Walker (YAT ICT Manager)"),
        ("Status", "As built — this is the deployed environment"),
        ("Classification", "Internal — YAT ICT, Student Administration, and MTS personnel on signed MSA"),
    ])
    contents(doc)

    h1 = lambda t: doc.add_paragraph(t, style="Heading 1")
    h3 = lambda t: doc.add_paragraph(t, style="Heading 3")

    h1("1. Purpose and Scope")
    add_body_paragraph(doc, "This document specifies the baseline AWS architecture MTS implemented as the foundation "
                 "build phase of the YAT Enrolline cloud migration, for the student records and enrolment "
                 "management system. The design stops at “infrastructure ready for application deployment”: "
                 "the EC2 instance is provisioned with the OS, the RDS instance with an empty PostgreSQL "
                 "engine, the S3 bucket empty, and the load balancer with placeholder health checks — no "
                 "application binaries or student data were placed by the MTS build.")
    add_body_paragraph(doc, "It is the **as-built record of the current environment**, and the starting point for "
                 "any onward improvement of that environment.")
    h3("In scope of this design")
    add_bullet_list(doc, [
        "The production cloud foundation for the Enrolline student records application.",
        "All compute, networking, identity, storage, database, autoscaling and baseline monitoring needed to run Enrolline as a staff-facing workload in AWS.",
        "An Amazon S3 document store for scanned student attachments, which are a working part of the record rather than an archive.",
        "Single-region, single-Availability-Zone deployment in ap-southeast-2 (Sydney).",
    ])
    lab_note(doc, "the design calls for ap-southeast-2 (Sydney). If you are building this in an AWS "
                  "Academy Learner Lab, deploy to whichever region that environment gives you — us-east-1 — "
                  "and treat it as standing in for Sydney throughout. That is acceptable.")
    h3("Out of scope — deferred to the follow-on HA design phase")
    add_bullet_list(doc, [
        "High-availability hardening (Multi-AZ database, cross-AZ compute resilience, failure-simulation testing).",
        "Disaster recovery to a second AWS region; DR runbook and tabletop testing.",
        "Capacity scheduling against the academic calendar.",
        "Application re-platforming (Enrolline remains Amazon Linux + PostgreSQL).",
    ])
    h3("Out of MTS scope entirely — YAT ICT responsibility")
    add_bullet_list(doc, [
        "Enrolline application installation onto the EC2 instance(s) after handover.",
        "Database migration (extract from the on-premises database, load into RDS for PostgreSQL).",
        "Migration of the scanned document library into the S3 store.",
        "Cutover — DNS switch, parallel running, decommissioning, user redirection (avoiding intake and census windows).",
        "Organisational change management — CAB approvals, communications, training, post-cutover support.",
    ])

    h1("2. Design Inputs and Requirements")
    h3("2.1 Inputs")
    add_bullet_list(doc, [
        "Enrolline Application Specification — functions, user load, data, integrations, SLAs, data residency.",
        "Enrolline Infrastructure Specifications — current operational state, utilisation, growth.",
        "Engagement Role Brief and Consultation Notes — engagement scope; OS, application and database preservation.",
        "YAT academic calendar — the intake enrolment windows and census dates that fix the Restricted Period.",
    ])
    h3("2.2 Requirements the design must meet")
    add_data_table(doc, ["Requirement", "Target / note"], [
        ["Availability (business hours, outside intake)", "99.5%"],
        ["Availability (intake enrolment window)", "99.9% expected by the business"],
        ["RPO", "≤ 1 hour"],
        ["RTO", "≤ 2 hours"],
        ["Concurrency (typical / intake peak)", "25–40 / 90–110"],
        ["Data residency", "Australian region only — Privacy Act 1988, APP 8"],
        ["Student-record retention", "30 years (NVR Standards / ASQA)"],
        ["Application preservation", "No change to Enrolline; PostgreSQL retained"],
    ], widths=[7.5, 9.0])

    h1("3. Review of Existing Architecture")
    add_body_paragraph(doc, "Enrolline previously ran on a single on-premises server co-locating the application and "
                 "its database, with scanned attachments on a file share. Its single points of failure were "
                 "the server itself, its storage, and the campus power and network feeding it. Recovery "
                 "depended on nightly tape and was never tested against the intake calendar.")

    h1("4. Architecture Design")
    h3("4.1 Assumptions and constraints")
    add_data_table(doc, ["#", "Assumption / constraint", "Source"], [
        ["A1", "Enrolline runs unchanged on Amazon Linux 2023 with PostgreSQL", "Role Brief"],
        ["A2", "Students do not access Enrolline; it is staff-facing only", "Application Specification"],
        ["A3", "Intake and census dates are fixed externally and published in advance", "Registrar"],
        ["A4", "Scanned attachments must be retained for 30 years", "NVR Standards"],
        ["C1", "Single-AZ for the foundation build; HA deferred", "Engagement scope"],
        ["C2", "All data remains in ap-southeast-2", "Privacy / Data Handling Policy"],
    ], widths=[1.5, 10.0, 5.0])

    h3("4.2 AWS account and region")
    add_body_paragraph(doc, "A dedicated AWS account in region ap-southeast-2 (Sydney). No resources are created "
                 "outside that region.")

    h3("4.3 Identity and Access Management (IAM)")
    add_data_table(doc, ["Group", "Purpose", "Indicative permissions"], [
        ["EnrollineAdmins", "YAT ICT administration of the environment", "EC2, RDS, S3, CloudWatch — full within the account"],
        ["EnrollineReadOnly", "Audit and review access", "ReadOnlyAccess"],
        ["EnrollineAppRole", "Instance profile for the application tier", "Read/write to yat-enrolline-documents; CloudWatch Logs write"],
    ], widths=[4.5, 5.5, 6.5])
    lab_note(doc, "the Learner Lab does not permit IAM user, group or role creation. Use the role the "
                  "lab provides and note where these groups would sit in a real account.")

    h3("4.4 Network topology")
    add_data_table(doc, ["Subnet", "CIDR", "Tier", "Internet-facing?"], [
        ["enrolline-public-a", "10.30.1.0/24", "Public (ALB, NAT)", "Yes"],
        ["enrolline-public-b", "10.30.2.0/24", "Public (ALB only)", "Yes"],
        ["enrolline-app-a", "10.30.11.0/24", "Application", "No — egress via NAT"],
        ["enrolline-data-a", "10.30.21.0/24", "Database", "No"],
        ["enrolline-data-b", "10.30.22.0/24", "Database (subnet group only)", "No"],
    ], widths=[4.5, 3.0, 5.5, 3.5])
    add_body_paragraph(doc, "VPC enrolline-vpc, 10.30.0.0/16. A single NAT Gateway in enrolline-public-a provides "
                 "outbound internet for the application tier. There is no application subnet in the second "
                 "Availability Zone. There are no VPC endpoints, so traffic to Amazon S3 and to the USI "
                 "Registry web service leaves the VPC through the NAT Gateway.")
    diagram_figure(doc,
                   "Figure 4.4 — Enrolline baseline network topology. The workload runs in "
                   "ap-southeast-2a; the second-zone subnets carry only the load balancer and the "
                   "database subnet group. Amazon S3 is a regional service and sits outside the "
                   "VPC, reached through the NAT Gateway.",
                   "network-enrolline-baseline-singleaz.png")

    h3("4.5 Compute (EC2 + Auto Scaling)")
    add_data_table(doc, ["Attribute", "Value"], [
        ["Instance family", "General-purpose burstable (t3.micro / t3.small)"],
        ["AMI", "Amazon Linux 2023 + Enrolline Student Management Suite"],
        ["Auto Scaling Group", "min 1 / desired 1 / max 2 — enrolline-app-a only"],
        ["Scaling policy", "Target tracking on CPU at 70%"],
        ["Capacity basis", "Sized for the intake peak and held at that size year-round"],
        ["Administrative access", "Systems Manager Session Manager — no key pair, no bastion"],
    ], widths=[5.5, 11.0])

    h3("4.6 Load balancing (ALB)")
    add_body_paragraph(doc, "An internet-facing Application Load Balancer, enrolline-alb, spanning enrolline-public-a "
                 "and enrolline-public-b, with an HTTP:80 listener to target group enrolline-tg. Health "
                 "check HTTP GET on /, 30-second interval, 2 failures to remove a target. TLS is not "
                 "terminated at this baseline.")

    h3("4.7 Database (RDS for PostgreSQL)")
    add_data_table(doc, ["Attribute", "Value"], [
        ["Engine", "Amazon RDS for PostgreSQL"],
        ["Instance class", "db.t3.micro / db.t3.small"],
        ["Multi-AZ", "Disabled — no standby instance"],
        ["Storage", "gp3, 20 GB; encrypted with AWS KMS"],
        ["Placement", "enrolline-data-a; not publicly accessible"],
        ["Subnet group", "enrolline-db-subnet-group, spanning both data subnets"],
        ["Backup retention", "7 days, automated, with transaction logs"],
    ], widths=[5.5, 11.0])

    h3("4.8 Storage — the S3 document store")
    add_data_table(doc, ["Attribute", "Value"], [
        ["Bucket", "yat-enrolline-documents"],
        ["Contents", "Scanned ID evidence, prior qualifications, USI evidence, support plans"],
        ["Versioning", "Enabled"],
        ["Encryption", "SSE-S3"],
        ["Storage class", "S3 Standard for all objects — no lifecycle configuration"],
        ["Public access", "Blocked at the bucket level"],
        ["Access path", "Via the NAT Gateway — no S3 VPC endpoint"],
    ], widths=[5.5, 11.0])

    h3("4.9 Security")
    add_data_table(doc, ["Security group", "Inbound"], [
        ["enrolline-alb-sg", "HTTP:80 from 0.0.0.0/0"],
        ["enrolline-app-sg", "HTTP:80 from enrolline-alb-sg only"],
        ["enrolline-db-sg", "PostgreSQL:5432 from enrolline-app-sg only"],
    ], widths=[5.5, 11.0])
    add_body_paragraph(doc, "Encryption at rest is enabled on the database and the document store. No VPC flow logs, "
                 "load-balancer access logs, or S3 access logging are configured at this baseline.")

    h3("4.10 Monitoring (baseline)")
    add_bullet_list(doc, [
        "enrolline-unhealthy-hosts — any unhealthy target behind the load balancer.",
        "enrolline-db-storage-low — free database storage below 15% of allocation.",
    ])

    h3("4.11 Naming and tagging conventions")
    add_body_paragraph(doc, "All resources are prefixed enrolline- and tagged System=Enrolline, "
                 "Environment=Production, Owner=YAT-ICT.")

    h3("4.12 Backup")
    add_bullet_list(doc, [
        "Database: RDS automated daily backups, 7-day retention, plus transaction logs.",
        "Documents: S3 versioning. No cross-Region copy.",
    ])

    h3("4.13 Recovery objectives — baseline state")
    add_body_paragraph(doc, "The two-hour recovery-time objective is **not reliably met** by this baseline. "
                 "Recovering a failed single-AZ database depends on point-in-time restore, which does not "
                 "complete inside two hours at this data volume. This is a known limitation of the "
                 "foundation build, recorded here rather than resolved.")

    h3("4.14 Single points of failure")
    add_bullet_list(doc, [
        "RDS database — single instance, one Availability Zone, no standby.",
        "EC2 application — all capacity in enrolline-app-a, with no second-zone application subnet to expand into.",
        "NAT Gateway — single, in enrolline-public-a. Its loss also removes S3 document access and USI verification.",
    ])

    h3("4.15 Configuration decisions left to the implementer")
    add_bullet_list(doc, [
        "Instance and database class within the stated families, on the documented load profile.",
        "CloudWatch alarm thresholds beyond the two baseline alarms.",
        "Tag values beyond the three mandatory keys.",
    ])

    h1("5. Implementation Sequencing")
    add_bullet_list(doc, [
        "Network — VPC, subnets, gateways, routing, security groups.",
        "Storage — the S3 document store and its bucket policy.",
        "Database — subnet group, parameter group, RDS instance.",
        "Compute — launch template, Auto Scaling group, target group, load balancer.",
        "Monitoring — the two baseline alarms.",
    ])

    h1("6. Verification Plan")
    add_bullet_list(doc, [
        "The load balancer reports a healthy target and serves the application endpoint.",
        "The application tier can read and write the document store.",
        "The database is reachable from the application tier and from nowhere else.",
        "Both baseline alarms exist and are in an OK state.",
    ])

    h1("7. References")
    add_bullet_list(doc, [
        "Enrolline Application Specification (YAT ICT records)",
        "Enrolline Infrastructure Specifications (YAT ICT records)",
        "Enrolline Network Diagram (YAT ICT records)",
        "Enrolline Operational Costing (YAT ICT records)",
        "Backup and Retention Policy; Change Management Procedure; Privacy / Data Handling Policy (intranet)",
    ])

    document_control(doc, h1, "v1.0 — as built",
                     "Pat Lin (MTS Senior Consultant) · Sam Walker (YAT ICT Manager)")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# --------------------------------------------------------------------------------------
# IMPROVED — the approved "to be" design
# --------------------------------------------------------------------------------------

def build_improved(path):
    doc = Document()
    configure_styles(doc)
    page_setup(doc)

    cover(doc, "YAT Enrolline Cloud Architecture — Approved Improvement Design", [
        ("Engagement", "YAT Enrolline Cloud Infrastructure Improvement"),
        ("Document type", "Technical design (Solution Design)"),
        ("Version", "v1.0 — Approved for implementation"),
        ("Authored by", "MTS Improvement Team, supervised by Pat Lin (MTS Senior Consultant)"),
        ("Approved by", "Sam Walker (YAT ICT Manager) · Jess Tran (YAT Registrar) · Asha Rao (YAT Compliance)"),
        ("Supersedes", "Nothing — it improves the Baseline Design, which remains the as-was record"),
        ("Classification", "Internal — YAT ICT, Student Administration, Compliance, and MTS personnel on signed MSA"),
    ])
    contents(doc)

    h1 = lambda t: doc.add_paragraph(t, style="Heading 1")
    h3 = lambda t: doc.add_paragraph(t, style="Heading 3")

    h1("1. Purpose and Scope")
    add_body_paragraph(doc, "This document specifies the improvements YAT approved at the Phase 1 design review, and "
                 "is the design the Improvement Team implements in Phase 2. It takes the single-Availability-"
                 "Zone Enrolline baseline and improves it across security, reliability, scalability and cost, "
                 "plus a light India residency slice for regulatory logs.")
    add_body_paragraph(doc, "The work is divided across four components — network, compute, database and storage — "
                 "one per Improvement Team member. The Enrolline application and its student records are "
                 "preserved unchanged throughout.")

    h1("2. Approved Improvements by Component")

    h3("2.1 Network")
    add_data_table(doc, ["Change", "Why"], [
        ["Add application subnet enrolline-app-b in ap-southeast-2b", "Gives the compute tier somewhere to place second-zone capacity"],
        ["Add a Gateway VPC endpoint for Amazon S3", "Document traffic stops leaving the VPC through the NAT Gateway — removes a data-processing charge and a dependency"],
        ["Enable VPC flow logs to CloudWatch Logs", "Required to detect and investigate a personal-data breach; no network record exists today"],
        ["Enable load-balancer access logs to S3", "Access record for the staff-facing entry point"],
    ], widths=[7.5, 9.0])

    h3("2.2 Compute")
    add_data_table(doc, ["Change", "Why"], [
        ["Extend the Auto Scaling group across enrolline-app-a and enrolline-app-b", "Removes the single-zone application failure mode"],
        ["Raise max capacity to 4", "Headroom for the intake peak once capacity is no longer held constantly"],
        ["Add scheduled scaling against the academic calendar", "The peaks are published months ahead; holding peak capacity for forty-six quiet weeks is the largest identified inefficiency"],
    ], widths=[7.5, 9.0])
    add_body_paragraph(doc, "Scheduled scaling raises desired capacity ahead of each intake enrolment window and each "
                 "census date, and returns it afterwards. The schedule is derived from the published academic "
                 "calendar, not from observed load, so capacity is in place before the load arrives.")

    h3("2.3 Database")
    add_data_table(doc, ["Change", "Why"], [
        ["Convert the RDS instance to a Multi-AZ deployment with an automatic-failover standby", "Brings recovery inside the two-hour objective, which point-in-time restore does not reliably meet"],
    ], widths=[7.5, 9.0])
    add_body_paragraph(doc, "Multi-AZ is a synchronous standby in the second Availability Zone with automatic "
                 "failover. It is not a read replica and adds no read capacity; it is bought for recovery time, "
                 "and its cost is approximately a doubling of the database line.")
    add_body_paragraph(doc, "This is the **only** database property the improvement changes. The change is applied "
                 "in place, as a modification to the existing instance — nothing in this design replaces the "
                 "database or moves student records, and no other database property is touched.")
    lab_note(doc, "the lab role may refuse rds:ModifyDBInstance. Where it does, the Multi-AZ conversion is "
                  "specified and justified in the design and verified by inspection of the template rather "
                  "than by a live change. Say which you did.")

    h3("2.4 Storage")
    add_data_table(doc, ["Change", "Why"], [
        ["Add an S3 lifecycle configuration — Standard → Standard-IA at 90 days → Glacier Instant Retrieval at 2 years", "Attachment reads fall away once an enrolment closes, but 30-year retention means nothing is ever deleted"],
        ["Enable S3 server access logging", "Access record over student identity evidence"],
        ["Confirm Block Public Access and versioning remain enforced", "No change intended; verified as part of the improvement"],
    ], widths=[7.5, 9.0])
    add_body_paragraph(doc, "The lifecycle configuration does not delete anything. Retention is unchanged at 30 years; "
                 "only the storage class changes as objects age.")

    h1("3. India Residency Slice")
    add_body_paragraph(doc, "The CERT-In Directions require system and access logs for the India-related operation to "
                 "be held within Indian jurisdiction for a rolling 180 days and to be retrievable quickly "
                 "enough to support a six-hour incident-reporting obligation.")
    add_bullet_list(doc, [
        "System and access logs for the India operation are replicated to a log destination in the Mumbai region (ap-south-1) with 180-day retention.",
        "The main system, the student records and the document store remain in Sydney — the DPDP Act permits this, and the Australian residency position is unchanged.",
        "Only logs cross the border. No student personal data is moved to India by this design.",
    ])
    na(doc, "no deployable lab variant is produced for the residency slice; it is specified at design level and "
            "assessed as a design decision.")

    h1("4. Goals and Performance Metrics")
    add_data_table(doc, ["Goal", "Metric", "Target"], [
        ["Reliability", "Recovery time from database failure", "≤ 2 hours, evidenced by failover"],
        ["Reliability", "Availability during an intake window", "99.9%"],
        ["Scalability", "Capacity available at the start of an intake window", "In place before the window opens, by schedule"],
        ["Security", "Network and access logging coverage", "Flow logs, ALB access logs, S3 access logs all enabled"],
        ["Cost", "Compute cost outside intake windows", "Reduced against the year-round peak-capacity baseline"],
    ], widths=[3.5, 7.0, 6.0])

    h1("5. What Is Deliberately Not Changed")
    add_bullet_list(doc, [
        "The Enrolline application and its PostgreSQL engine — IR-4.",
        "The Sydney region for all student records and documents.",
        "The commercial Enrolline licensing and vendor support, which are fixed against the user population and unaffected by infrastructure change.",
        "Retention periods for any record class.",
    ])

    h1("6. Implementation and Change Control")
    add_body_paragraph(doc, "All production-affecting work is applied as a change-set against the deployed baseline "
                 "and scheduled outside the Restricted Period — the intake enrolment windows, the week either "
                 "side of a census date, and the January AVETMISS submission window. Every change in this "
                 "design is additive or in-place; none replaces the database or moves data.")

    h1("7. References")
    add_bullet_list(doc, [
        "Enrolline Cloud Architecture — Baseline Design — the as-was state this design improves",
        "Improvement Requirements (IR-1…7) — the outcomes this design is justified against",
        "Indian Regulatory Requirements — the obligations behind §3",
        "Enrolline Operational Costing — the cost base each improvement is measured against",
        "Change Management Procedure (intranet policies)",
    ])

    document_control(doc, h1, "v1.0 — approved for implementation",
                     "Sam Walker (YAT ICT Manager) · Jess Tran (YAT Registrar) · Asha Rao (YAT Compliance)")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--only", choices=["baseline", "improved"],
                    help="build just one of the two documents")
    args = ap.parse_args()

    if args.only in (None, "baseline"):
        p = build_baseline(OUT_DIR / "YAT-Enrolline-Baseline-Solution-Design.docx")
        print(f"Wrote {p}")
    if args.only in (None, "improved"):
        p = build_improved(OUT_DIR / "YAT-Enrolline-Improved-Solution-Design.docx")
        print(f"Wrote {p}")


if __name__ == "__main__":
    main()
