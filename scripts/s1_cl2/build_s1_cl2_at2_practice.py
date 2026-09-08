#!/usr/bin/env python3
"""Build the S1-CL2 AT2 PRACTICE microservice-and-IaC run sheet (.docx).

A practice exercise, not an assessment: a branded YAT/MTS document with no Kangan wrapper, no
marking criteria and no UoC tags. Same shape as the AT2 assessment workbook, on the LMS
activity-audit service rather than the website's access-log service — different template,
different fault, different field names — and with the click-by-click detail the assessment
withholds.

It renders through the assessment's own renderer (s1_cl2_at2_run_sheet.render, and its
render_supplied for the two provided files) with its content lists passed in, so the practice
and the assessment cannot drift structurally even though every value in them differs.

Usage:  python scripts/s1_cl2/build_s1_cl2_at2_practice.py [output.docx]
Default: S1-CL2-Cloud-Disaster-Recovery/delivery/practice/AT2-Practice-Microservice-IaC-Run-Sheet.docx
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # noqa: E402
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # noqa: E402
sys.path.insert(0, str(next(d / "scripts" for d in Path(__file__).resolve().parents
                            if (d / "scripts" / "helpers" / "__init__.py").exists())))  # noqa: E402

import s1_cl2_at2_practice_run_sheet as content  # noqa: E402
import s1_cl2_at2_run_sheet as at2  # noqa: E402
from helpers import run_sheet as R  # noqa: E402
from helpers.docx_styling import paragraph_bottom_rule  # noqa: E402
from helpers.scenario_document import build_header_footer, configure_styles, wordmark  # noqa: E402
from brand import ADDRESS, GREY, TEAL  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.shared import Pt, Cm, RGBColor  # noqa: E402


def build(path):
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.6); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)
    build_header_footer(sec)

    # ---- cover ----
    wordmark(doc.add_paragraph())
    ar = doc.add_paragraph().add_run(ADDRESS)
    ar.font.size = Pt(9); ar.font.color.rgb = RGBColor.from_string(GREY)
    paragraph_bottom_rule(doc.add_paragraph(), TEAL, sz=12)
    for _ in range(3):
        doc.add_paragraph()
    doc.add_paragraph(style="Title").add_run("Practice — Microservice & IaC Build")
    sub = doc.add_paragraph().add_run("LMS activity-audit service — build run sheet")
    sub.font.size = Pt(15); sub.italic = True; sub.font.color.rgb = RGBColor.from_string(GREY)

    doc.add_section(WD_SECTION.NEW_PAGE); build_header_footer(doc.sections[-1])
    h1 = lambda t: doc.add_paragraph(t, style="Heading 1")
    h2 = lambda t: R.heading2(doc, t)

    h1("The engagement")
    for para in content.SCENARIO:
        R.p(doc, para, after=8)
    R.resources_block(doc, content.RESOURCES)
    h1("How to work through this")
    for para in content.INSTRUCTIONS:
        R.p(doc, para, after=8)
    R.note(doc, content.REGION_NOTE)

    at2.render_supplied(doc, h1, h2,
                        intro=content.SUPPLIED_INTRO,
                        datastore=content.DATASTORE_YAML,
                        handler=content.HANDLER_SUMMARY,
                        contract=content.WEBHOOK_CONTRACT,
                        names=("lms-activity-store.yaml", "activity_writer.py"),
                        producer="LMS")

    at2.render(doc, h1, h2, mode="student",
               tasks=content.TASKS, questions=[], notes=True)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"Wrote {path}")
    print(f"  {len(content.TASKS)} tasks")


if __name__ == "__main__":
    default = ("S1-CL2-Cloud-Disaster-Recovery/delivery/practice/"
               "AT2-Practice-Microservice-IaC-Run-Sheet.docx")
    build(sys.argv[1] if len(sys.argv) > 1 else default)
