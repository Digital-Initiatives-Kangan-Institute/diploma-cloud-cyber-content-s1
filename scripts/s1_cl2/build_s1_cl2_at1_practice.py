#!/usr/bin/env python3
"""Build the S1-CL2 AT1 PRACTICE design + DR-plan run sheet (.docx).

A practice exercise, not an assessment: a branded YAT/MTS document with no Kangan wrapper, no
marking criteria and no UoC tags. Same three-part shape as the AT1 assessment workbook — design,
DR plan, approval — on the LMS Global Expansion engagement rather than the website one, and with
the teaching the assessment withholds: "Things to consider" leading questions on every task, one
worked exemplar row per table, click-by-click steps where a task is procedural, and a personal
notes box throughout.

It renders through the assessment's own renderers (s1_cl2_at1_part_{a,b,c}_run_sheet.render)
with its content lists passed in, so the practice and the assessment cannot drift structurally
even though every value in them differs.

Usage:  python scripts/s1_cl2/build_s1_cl2_at1_practice.py [output.docx]
Default: S1-CL2-Cloud-Disaster-Recovery/delivery/practice/AT1-Practice-Design-DR-Plan-Run-Sheet.docx
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # noqa: E402
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # noqa: E402
sys.path.insert(0, str(next(d / "scripts" for d in Path(__file__).resolve().parents
                            if (d / "scripts" / "helpers" / "__init__.py").exists())))  # noqa: E402

import s1_cl2_at1_part_a_practice_run_sheet as prac_a  # noqa: E402
import s1_cl2_at1_part_b_practice_run_sheet as prac_b  # noqa: E402
import s1_cl2_at1_part_c_practice_run_sheet as prac_c  # noqa: E402
import s1_cl2_at1_part_a_run_sheet as part_a  # noqa: E402
import s1_cl2_at1_part_b_run_sheet as part_b  # noqa: E402
import s1_cl2_at1_part_c_run_sheet as part_c  # noqa: E402
from helpers import run_sheet as R  # noqa: E402
from helpers.docx_styling import paragraph_bottom_rule  # noqa: E402
from helpers.scenario_document import build_header_footer, configure_styles, wordmark  # noqa: E402
from brand import ADDRESS, GREY, TEAL  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.shared import Pt, Cm, RGBColor  # noqa: E402


def build(path):
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.6); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)
    build_header_footer(sec)

    # ---- cover ----
    wordmark(doc.add_paragraph())
    ar = doc.add_paragraph().add_run(ADDRESS)
    ar.font.size = Pt(9); ar.font.color.rgb = RGBColor.from_string(GREY)
    paragraph_bottom_rule(doc.add_paragraph(), TEAL, sz=12)
    for _ in range(3):
        doc.add_paragraph()
    doc.add_paragraph(style="Title").add_run("Practice — Design & DR Plan")
    sub = doc.add_paragraph().add_run("LMS Global Expansion — design, recovery plan and approval")
    sub.font.size = Pt(15); sub.italic = True; sub.font.color.rgb = RGBColor.from_string(GREY)

    doc.add_section(WD_SECTION.NEW_PAGE); build_header_footer(doc.sections[-1])
    h1 = lambda t: doc.add_paragraph(t, style="Heading 1")
    h2 = lambda t: R.heading2(doc, t)

    h1("The engagement")
    for para in prac_a.SCENARIO:
        R.p(doc, para, after=8)
    h1("How to work through this")
    for para in prac_a.INSTRUCTIONS:
        R.p(doc, para, after=8)

    part_a.render(doc, h1, h2, mode="student",
                  design=prac_a.DESIGN, questions=[],
                  current_arch=prac_a.CURRENT_ARCH,
                  network_diagram=prac_a.NETWORK_DIAGRAM,
                  scope_note=prac_a.SCOPE_NOTE, notes=True)

    part_b.render(doc, h1, h2, mode="student",
                  plan=prac_b.PLAN, questions=[],
                  intro=prac_b.INTRO, resources=prac_b.RESOURCES,
                  template_note=prac_b.TEMPLATE_NOTE, notes=True)

    part_c.render(doc, h1, h2, mode="student",
                  approval=prac_c.APPROVAL, intro=prac_c.INTRO,
                  event=prac_c.EVENT, resources=prac_c.RESOURCES, notes=True)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"Wrote {path}")
    print(f"  Part A {len(prac_a.DESIGN)} tasks  ·  Part B {len(prac_b.PLAN)} tasks  ·  "
          f"Part C {len(prac_c.APPROVAL)} tasks")


if __name__ == "__main__":
    default = ("S1-CL2-Cloud-Disaster-Recovery/delivery/practice/"
               "AT1-Practice-Design-DR-Plan-Run-Sheet.docx")
    build(sys.argv[1] if len(sys.argv) > 1 else default)
