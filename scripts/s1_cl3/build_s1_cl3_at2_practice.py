#!/usr/bin/env python3
"""Build the S1-CL3 AT2 PRACTICE run sheet (.docx).

A practice exercise, not an assessment: a branded YAT/MTS document with no Kangan wrapper, no
marking criteria and no UoC tags. Same shape as the AT2 assessment workbook, on the public website improvement build rather
than the Ledgerline one — and with the teaching the assessment withholds.

It renders through the assessment's own renderer (s1_cl3_at2_run_sheet.render) with its content
lists passed in, so the practice and the assessment cannot drift structurally even though every
value in them differs.

Usage:  python scripts/s1_cl3/build_s1_cl3_at2_practice.py [output.docx]
Default: S1-CL3-Cloud-Infrastructure-Improvement/delivery/practice/AT2-Practice-Team-Implementation-Run-Sheet.docx
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # noqa: E402
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # noqa: E402
sys.path.insert(0, str(next(d / "scripts" for d in Path(__file__).resolve().parents
                            if (d / "scripts" / "helpers" / "__init__.py").exists())))  # noqa: E402

import s1_cl3_at2_practice_run_sheet as content  # noqa: E402
import s1_cl3_at2_run_sheet as at  # noqa: E402
from helpers import run_sheet as R  # noqa: E402
from helpers.docx_styling import paragraph_bottom_rule  # noqa: E402
from helpers.scenario_document import build_header_footer, configure_styles, wordmark  # noqa: E402
from brand import ADDRESS, GREY, TEAL  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.shared import Pt, Cm, RGBColor  # noqa: E402


def build(path):
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.6); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)
    build_header_footer(sec)

    wordmark(doc.add_paragraph())
    ar = doc.add_paragraph().add_run(ADDRESS)
    ar.font.size = Pt(9); ar.font.color.rgb = RGBColor.from_string(GREY)
    paragraph_bottom_rule(doc.add_paragraph(), TEAL, sz=12)
    for _ in range(3):
        doc.add_paragraph()
    doc.add_paragraph(style="Title").add_run("Practice — Team Implementation")
    sub = doc.add_paragraph().add_run("YAT public website — working and leading in a team")
    sub.font.size = Pt(15); sub.italic = True; sub.font.color.rgb = RGBColor.from_string(GREY)

    doc.add_section(WD_SECTION.NEW_PAGE); build_header_footer(doc.sections[-1])
    h1 = lambda t: doc.add_paragraph(t, style="Heading 1")
    h2 = lambda t: R.heading2(doc, t)

    h1("The engagement")
    for para in content.SCENARIO:
        R.p(doc, para, after=8)
    h1("How this exercise works")
    R.settings_table(doc, content.HOW_ASSESSED)
    h1("How to work through this")
    for para in content.INSTRUCTIONS:
        R.p(doc, para, after=8)

    at.render(doc, h1, h2, mode="student",
              plan=content.PLAN, work=content.WORK, meeting=content.MEETING,
              running=content.RUNNING, review=content.REVIEW, questions=[],
              meeting_intro="Arrange this early — with your teacher, or with a classmate "
                            "from another team. Nothing is signed off for real; what you "
                            "want is an outside opinion while there is still time to act "
                            "on it.",
              notes=True)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"Wrote {path}")
    print(f"  {len(content.PLAN)+len(content.WORK)+len(content.MEETING)
            +len(content.RUNNING)+len(content.REVIEW)} tasks across five parts")


if __name__ == "__main__":
    default = ("S1-CL3-Cloud-Infrastructure-Improvement/delivery/practice/AT2-Practice-Team-Implementation-Run-Sheet.docx")
    build(sys.argv[1] if len(sys.argv) > 1 else default)
